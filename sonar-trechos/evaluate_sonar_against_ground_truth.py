
import csv
import json
import math
from pathlib import Path

from openpyxl import load_workbook


GROUND_TRUTH_XLSX = Path("Analise-base-MLCQ/MLCQ_ground_truth_completo.xlsx")
SONAR_PREDICTIONS_JSONL = Path("sonar-trechos/sonar_sample_predictions.jsonl")

OUTPUT_DIR = Path("sonar-trechos/evaluation")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_COMPARISON_CSV = OUTPUT_DIR / "comparison_sample_smell.csv"
OUTPUT_METRICS_CSV = OUTPUT_DIR / "metrics_by_smell.csv"
OUTPUT_SUMMARY_TXT = OUTPUT_DIR / "evaluation_summary.txt"
OUTPUT_SONAR_ONLY_JSONL = OUTPUT_DIR / "sonar_only_without_ground_truth.jsonl"


def normalize_text(value):
    if value is None:
        return ""

    text = str(value).strip().lower()

    replacements = {
        "_": " ",
        "-": " ",
        "/": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return " ".join(text.split())


def normalize_smell(value):
    text = normalize_text(value)

    mapping = {
        "blob": "blob",
        "god class": "blob",
        "blob god class": "blob",
        "blob class": "blob",

        "data class": "data-class",
        "dataclass": "data-class",

        "feature envy": "feature-envy",
        "featureenvy": "feature-envy",

        "long method": "long-method",
        "longmethod": "long-method",
        "long method statistical": "long-method",
    }

    return mapping.get(text, text.replace(" ", "-"))


def parse_verdict(value):
    """
    Retorna:
    1 para tem smell
    0 para não tem smell
    None para indefinido / vazio / não avaliado
    """

    if value is None:
        return None

    if isinstance(value, bool):
        return 1 if value else 0

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isnan(value) if isinstance(value, float) else False:
            return None

        if int(value) == 1:
            return 1

        if int(value) == 0:
            return 0

        if int(value) == -1:
            return None

    text = normalize_text(value)

    positives = {
        "1", "sim", "yes", "true", "verdadeiro", "positivo",
        "tem", "has", "smell", "presente",
        "tem smell", "tem_smell"
    }

    negatives = {
        "0", "nao", "não", "no", "false", "falso", "negativo",
        "nao tem", "não tem", "ausente", "none",
        "nao tem smell", "não tem smell", "nao_tem_smell"
    }

    unknowns = {
        "-1", "", "nan", "none", "null", "n/a", "na",
        "indefinido", "nao avaliado", "não avaliado"
    }

    if text in positives:
        return 1

    if text in negatives:
        return 0

    if text in unknowns:
        return None

    return None


def detect_columns(header):
    normalized = [normalize_text(cell) for cell in header]

    def find_column(possible_names, fallback_index):
        for name in possible_names:
            if name in normalized:
                return normalized.index(name)

        return fallback_index

    sample_col = find_column(["sample", "sample id", "sample_id"], 0)
    type_col = find_column(["type", "tipo"], 1)
    smell_col = find_column(["smell", "code smell", "target smell"], 2)

    # Coluna O = índice 14, pois começa em 0
    verdict_col = find_column(
        ["veredito", "verdict", "label", "target label", "ground truth"],
        14
    )

    return sample_col, type_col, smell_col, verdict_col


def load_ground_truth():
    wb = load_workbook(GROUND_TRUTH_XLSX, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        raise RuntimeError("Planilha de ground truth está vazia.")

    header = rows[0]
    sample_col, type_col, smell_col, verdict_col = detect_columns(header)

    ground_truth = {}
    raw_rows = []

    for row_number, row in enumerate(rows[1:], start=2):
        if not row or len(row) <= max(sample_col, smell_col, verdict_col):
            continue

        sample_raw = row[sample_col]
        smell_raw = row[smell_col]
        verdict_raw = row[verdict_col]
        type_raw = row[type_col] if len(row) > type_col else None

        if sample_raw is None or smell_raw is None:
            continue

        try:
            sample_id = int(sample_raw)
        except Exception:
            continue

        smell = normalize_smell(smell_raw)
        verdict = parse_verdict(verdict_raw)

        if verdict is None:
            continue

        key = (sample_id, smell)

        ground_truth[key] = {
            "sample_id": sample_id,
            "mlcq_type": type_raw,
            "smell": smell,
            "verdict": verdict,
            "original_verdict": verdict_raw,
            "row_number": row_number,
        }

        raw_rows.append(ground_truth[key])

    return ground_truth, raw_rows


def load_sonar_predictions():
    predictions = {}
    raw_predictions = []

    with SONAR_PREDICTIONS_JSONL.open("r", encoding="utf-8") as f:
        for line in f:
            if not line.strip():
                continue

            obj = json.loads(line)

            sample_id = int(obj["sample_id"])
            smell = normalize_smell(obj["detected_smell"])

            key = (sample_id, smell)

            if key not in predictions:
                predictions[key] = []

            predictions[key].append(obj)
            raw_predictions.append(obj)

    return predictions, raw_predictions


def classify(gt_verdict, sonar_detected):
    if gt_verdict == 1 and sonar_detected:
        return "TP"

    if gt_verdict == 1 and not sonar_detected:
        return "FN"

    if gt_verdict == 0 and sonar_detected:
        return "FP"

    if gt_verdict == 0 and not sonar_detected:
        return "TN"

    return "UNKNOWN"


def safe_div(num, den):
    return num / den if den else 0.0


def main():
    ground_truth, gt_rows = load_ground_truth()
    predictions, raw_predictions = load_sonar_predictions()

    comparison_rows = []
    metrics = {}

    all_smells = sorted(set([item["smell"] for item in gt_rows]))

    for gt in gt_rows:
        sample_id = gt["sample_id"]
        smell = gt["smell"]
        key = (sample_id, smell)

        sonar_detected = key in predictions
        result = classify(gt["verdict"], sonar_detected)

        if smell not in metrics:
            metrics[smell] = {"TP": 0, "FP": 0, "FN": 0, "TN": 0}

        metrics[smell][result] += 1

        comparison_rows.append({
            "sample_id": sample_id,
            "mlcq_type": gt["mlcq_type"],
            "smell": smell,
            "ground_truth": gt["verdict"],
            "original_verdict": gt["original_verdict"],
            "sonar_detected": int(sonar_detected),
            "result": result,
            "total_sonar_detections_for_pair": len(predictions.get(key, [])),
            "ground_truth_row": gt["row_number"],
        })

    sonar_only = []

    for key, detections in predictions.items():
        if key not in ground_truth:
            sample_id, smell = key
            sonar_only.append({
                "sample_id": sample_id,
                "smell": smell,
                "total_detections": len(detections),
                "detections": detections,
            })

    with OUTPUT_COMPARISON_CSV.open("w", encoding="utf-8", newline="") as f:
        fieldnames = [
            "sample_id",
            "mlcq_type",
            "smell",
            "ground_truth",
            "original_verdict",
            "sonar_detected",
            "result",
            "total_sonar_detections_for_pair",
            "ground_truth_row",
        ]

        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(comparison_rows)

    with OUTPUT_METRICS_CSV.open("w", encoding="utf-8", newline="") as f:
        fieldnames = [
            "smell",
            "TP",
            "FP",
            "FN",
            "TN",
            "precision",
            "recall",
            "f1",
            "accuracy",
        ]

        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()

        for smell in sorted(metrics):
            m = metrics[smell]

            tp = m["TP"]
            fp = m["FP"]
            fn = m["FN"]
            tn = m["TN"]

            precision = safe_div(tp, tp + fp)
            recall = safe_div(tp, tp + fn)
            f1 = safe_div(2 * precision * recall, precision + recall)
            accuracy = safe_div(tp + tn, tp + fp + fn + tn)

            writer.writerow({
                "smell": smell,
                "TP": tp,
                "FP": fp,
                "FN": fn,
                "TN": tn,
                "precision": round(precision, 4),
                "recall": round(recall, 4),
                "f1": round(f1, 4),
                "accuracy": round(accuracy, 4),
            })

    with OUTPUT_SONAR_ONLY_JSONL.open("w", encoding="utf-8") as f:
        for item in sonar_only:
            f.write(json.dumps(item, ensure_ascii=False) + "\n")

    total_tp = sum(m["TP"] for m in metrics.values())
    total_fp = sum(m["FP"] for m in metrics.values())
    total_fn = sum(m["FN"] for m in metrics.values())
    total_tn = sum(m["TN"] for m in metrics.values())

    micro_precision = safe_div(total_tp, total_tp + total_fp)
    micro_recall = safe_div(total_tp, total_tp + total_fn)
    micro_f1 = safe_div(2 * micro_precision * micro_recall, micro_precision + micro_recall)
    micro_accuracy = safe_div(
        total_tp + total_tn,
        total_tp + total_fp + total_fn + total_tn
    )

    with OUTPUT_SUMMARY_TXT.open("w", encoding="utf-8") as f:
        f.write("=== AVALIAÇÃO SONAR X MLCQ GROUND TRUTH ===\n\n")
        f.write(f"Linhas válidas do ground truth: {len(gt_rows)}\n")
        f.write(f"Detecções Sonar por ocorrência: {len(raw_predictions)}\n")
        f.write(f"Pares únicos sample+smell detectados pelo Sonar: {len(predictions)}\n")
        f.write(f"Detecções Sonar sem par correspondente no ground truth: {len(sonar_only)}\n\n")

        f.write("=== MÉTRICAS MICRO ===\n")
        f.write(f"TP: {total_tp}\n")
        f.write(f"FP: {total_fp}\n")
        f.write(f"FN: {total_fn}\n")
        f.write(f"TN: {total_tn}\n")
        f.write(f"Precision: {micro_precision:.4f}\n")
        f.write(f"Recall: {micro_recall:.4f}\n")
        f.write(f"F1: {micro_f1:.4f}\n")
        f.write(f"Accuracy: {micro_accuracy:.4f}\n\n")

        f.write("=== ARQUIVOS GERADOS ===\n")
        f.write(f"{OUTPUT_COMPARISON_CSV}\n")
        f.write(f"{OUTPUT_METRICS_CSV}\n")
        f.write(f"{OUTPUT_SONAR_ONLY_JSONL}\n")

    print("=== AVALIAÇÃO FINALIZADA ===")
    print(f"Linhas válidas do ground truth: {len(gt_rows)}")
    print(f"Detecções Sonar por ocorrência: {len(raw_predictions)}")
    print(f"Pares únicos sample+smell detectados pelo Sonar: {len(predictions)}")
    print(f"Detecções Sonar sem par correspondente no ground truth: {len(sonar_only)}")
    print()
    print(f"Resumo: {OUTPUT_SUMMARY_TXT}")
    print(f"Comparação: {OUTPUT_COMPARISON_CSV}")
    print(f"Métricas: {OUTPUT_METRICS_CSV}")
    print(f"Sonar sem GT: {OUTPUT_SONAR_ONLY_JSONL}")


if __name__ == "__main__":
    main()
